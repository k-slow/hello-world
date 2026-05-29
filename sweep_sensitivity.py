#!/usr/bin/env python3
"""
sweep_sensitivity.py  --  Single generator for the NDR sensitivity table (Tab 11).

The 3 engines compute ONE scenario at a time (whatever Tab 1 B136 resolves to).
To populate a full NDR curve (80/85/90/95/100%) we sweep: for each NDR level we
override every engine's per-deal NDR cell to that literal, recalc with LibreOffice,
and read the engine summary outputs. Results are written into Tab 11 as values plus
a generation stamp. The verifier re-runs this and asserts Tab 11 ties out.

Run:  python3 sweep_sensitivity.py            # writes results into Press_Model_v2.xlsx Tab 11
      python3 sweep_sensitivity.py --json      # just print the swept numbers as JSON
"""
import openpyxl, os, shutil, json, sys, datetime

MODEL = '/home/user/hello-world/Press_Model_v2.xlsx'
NDRS = [0.80, 0.85, 0.90, 0.95, 1.00]
# engine -> deal NDR rows (col G)
ENG_DEAL_ROWS = {'_Engine_NR': range(6, 11), '_Engine_R50': range(6, 15), '_Engine_R100': range(6, 15)}
# summary cells (identical across engines)
CELLS = {'10Y_DPI': 'B254', '15Y_DPI': 'B252', '10Y_IRR': 'B277', '15Y_IRR': 'B276',
         '10Y_LP': 'B253', '15Y_LP': 'B250'}


def recalc(src, outdir):
    home = outdir + '_home'
    os.system(f'rm -rf {outdir} {home} && mkdir -p {outdir} {home}')
    os.system(f'HOME={home} libreoffice --headless --calc --convert-to xlsx '
              f'--outdir {outdir} {src} >/dev/null 2>&1')
    return os.path.join(outdir, os.path.basename(src))


def sweep():
    results = {}
    for ndr in NDRS:
        tmp = f'/tmp/sweep_{int(ndr*100)}.xlsx'
        shutil.copy(MODEL, tmp)
        wb = openpyxl.load_workbook(tmp)
        for eng, rows in ENG_DEAL_ROWS.items():
            ws = wb[eng]
            for r in rows:
                ws.cell(r, 7).value = ndr           # col G = NDR, override scenario formula
        wb.save(tmp)
        calc = recalc(tmp, f'/tmp/sweep_out_{int(ndr*100)}')
        wv = openpyxl.load_workbook(calc, data_only=True)
        key = f'{int(ndr*100)}%'
        results[key] = {}
        for scen, eng in [('NR', '_Engine_NR'), ('R50', '_Engine_R50'), ('R100', '_Engine_R100')]:
            e = wv[eng]
            results[key][scen] = {m: e[c].value for m, c in CELLS.items()}
    return results


def write_tab11(results):
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.load_workbook(MODEL)
    ws = wb['11. NDR Sensitivity']
    # wipe everything
    for row in ws.iter_rows():
        for c in row:
            c.value = None
            c.font = Font()
            c.fill = PatternFill(fill_type=None)

    HF = Font(bold=True, color='FFFFFFFF'); HFILL = PatternFill('solid', fgColor='FF1F3864')
    SEC = Font(bold=True, color='FF1F3864'); SFILL = PatternFill('solid', fgColor='FFD9E1F2')
    BASEHL = PatternFill('solid', fgColor='FFFFF2CC')      # 85% base highlight
    BANDHL = PatternFill('solid', fgColor='FFE2EFDA')      # 80/90 band
    EXT = Font(italic=True, color='FF888888')              # 95/100 extended ref
    SUB = Font(italic=True, color='FF888888'); BOLD = Font(bold=True)

    ws['A1'] = 'FUND I  --  NDR SENSITIVITY (single source = engine sweep)'
    ws['A1'].font = HF; ws['A1'].fill = HFILL
    stamp = datetime.date.today().isoformat()
    ws['A2'] = (f'Generated {stamp} via sweep_sensitivity.py from the 3 engines. '
                f'Scenario-independent (full NDR curve). 80/85/90 = primary band; 95/100 = extended reference.')
    ws['A2'].font = SUB

    blocks = [('10Y DPI (LP MOIC)', '10Y_DPI', '0.00"x"', 4),
              ('15Y DPI (LP MOIC)', '15Y_DPI', '0.00"x"', 13),
              ('10Y STAGED IRR', '10Y_IRR', '0.0%', 22),
              ('15Y STAGED IRR', '15Y_IRR', '0.0%', 31)]
    for title, metric, numfmt, top in blocks:
        ws.cell(top, 1, title).font = SEC
        for c in range(1, 5):
            ws.cell(top, c).fill = SFILL
        hdr = top + 1
        ws.cell(hdr, 1, 'NDR').font = HF; ws.cell(hdr, 1).fill = HFILL
        for j, lab in enumerate(['No Recycle', '50% Recycle', '100% Recycle'], start=2):
            ws.cell(hdr, j, lab).font = HF; ws.cell(hdr, j).fill = HFILL
            ws.cell(hdr, j).alignment = Alignment(horizontal='center')
        for i, ndr in enumerate(NDRS):
            r = hdr + 1 + i
            key = f'{int(ndr*100)}%'
            ws.cell(r, 1, ndr).number_format = '0%'
            for j, scen in enumerate(['NR', 'R50', 'R100'], start=2):
                cell = ws.cell(r, j, results[key][scen][metric])
                cell.number_format = numfmt
                cell.alignment = Alignment(horizontal='center')
            # styling: highlight base 85, band 80/90, grey 95/100
            if abs(ndr - 0.85) < 1e-9:
                for c in range(1, 5):
                    ws.cell(r, c).fill = BASEHL; ws.cell(r, c).font = BOLD
            elif ndr in (0.80, 0.90):
                for c in range(1, 5):
                    ws.cell(r, c).fill = BANDHL
            else:
                for c in range(1, 5):
                    ws.cell(r, c).font = EXT

    note = 41
    ws.cell(note, 1, 'Base case = 85% NDR (highlighted). This is the curve; Tab 1 B136 selects which point the headline tabs feature.').font = SUB
    ws.cell(note+1, 1, 'Verification: 80% row == engine at Bear, 85% == Base, 90% == Bull (see verifier).').font = SUB
    ws.column_dimensions['A'].width = 10
    for c in 'BCD':
        ws.column_dimensions[c].width = 16
    wb.save(MODEL)


if __name__ == '__main__':
    res = sweep()
    if '--json' in sys.argv:
        print(json.dumps(res, indent=2, default=float))
    else:
        write_tab11(res)
        print('Tab 11 rebuilt from engine sweep.')
        for k in res:
            nr = res[k]['NR']
            print(f"  {k}: NR 10Y={nr['10Y_DPI']:.2f}x 15Y={nr['15Y_DPI']:.2f}x "
                  f"IRR10={nr['10Y_IRR']*100:.1f}% IRR15={nr['15Y_IRR']*100:.1f}%")
